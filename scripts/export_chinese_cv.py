#!/usr/bin/env python3
"""Render an accepted Chinese CV payload into the fixed Word format."""

import argparse
import json
import re
import tempfile
import zipfile
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


def font(run, size=9.25, bold=False, italic=False, name="微软雅黑"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def paragraph_base(p, before=0, after=0, line=1.0, keep=False):
    fmt = p.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing_rule = WD_LINE_SPACING.SINGLE
    fmt.line_spacing = line
    fmt.keep_together = keep


def bottom_rule(p):
    ppr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "14")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "000000")
    borders.append(bottom)
    ppr.append(borders)


def section_title(doc, text):
    p = doc.add_paragraph()
    paragraph_base(p, before=6, after=2, keep=True)
    font(p.add_run(text), size=11, bold=True)
    bottom_rule(p)


def role_line(doc, title, date_text):
    p = doc.add_paragraph()
    paragraph_base(p, before=4, keep=True)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(7.08), WD_TAB_ALIGNMENT.RIGHT)
    font(p.add_run(title), size=9.8, bold=True)
    font(p.add_run("\t" + date_text), size=9.2)


def organization_line(doc, text):
    if not text:
        return
    p = doc.add_paragraph()
    paragraph_base(p, keep=True)
    font(p.add_run(text), size=8.8, italic=True)


def bullet(doc, text):
    p = doc.add_paragraph()
    paragraph_base(p, after=1, line=1.02, keep=True)
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.first_line_indent = Inches(-0.17)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    font(p.add_run("•  "), bold=True)
    match = re.match(r"^(.+?[：:])(.+)$", text)
    if match:
        font(p.add_run(match.group(1)), bold=True)
        font(p.add_run(match.group(2)))
    else:
        font(p.add_run(text))


def entry(doc, item):
    role_line(doc, item["title"], item.get("date", ""))
    organization_line(doc, item.get("organization", ""))
    for text in item.get("bullets", []):
        bullet(doc, text)


def resolve_layout(payload):
    layout = payload.get("layout", {})
    profile = payload.get("profile", {})
    education = layout.get("education_position", "auto")
    projects = layout.get("projects_position", "auto")

    if education == "auto":
        graduation_year = profile.get("graduation_year")
        roles = int(profile.get("substantive_work_roles", len(payload.get("experience", []))))
        recent_graduate = isinstance(graduation_year, int) and date.today().year - graduation_year <= 2
        education = "before_experience" if recent_graduate or roles < 2 or profile.get("education_is_lead_signal") else "after_experience"

    if projects == "auto":
        projects = "before_experience" if profile.get("projects_are_primary_target_evidence") else "after_experience"

    allowed_education = {"before_experience", "after_experience"}
    allowed_projects = {"before_experience", "after_experience", "integrated"}
    if education not in allowed_education or projects not in allowed_projects:
        raise ValueError("Invalid layout route")
    return education, projects


def add_education(doc, items):
    if not items:
        return
    section_title(doc, "教育背景")
    for item in items:
        role_line(doc, item["title"], item.get("date", ""))


def add_entries(doc, title, items):
    if not items:
        return
    section_title(doc, title)
    for item in items:
        entry(doc, item)


def add_skills(doc, items):
    if not items:
        return
    section_title(doc, "专业认证与持续进修")
    for item in items:
        p = doc.add_paragraph()
        paragraph_base(p, keep=True)
        font(p.add_run(item["label"] + "："), size=9.3, bold=True)
        font(p.add_run(item["text"]), size=9.3)


def write_docx(payload, path):
    education_position, projects_position = resolve_layout(payload)
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.5)
    sec.bottom_margin = Inches(0.5)
    sec.left_margin = Inches(0.58)
    sec.right_margin = Inches(0.58)
    normal = doc.styles["Normal"]
    normal.font.name = "微软雅黑"
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "微软雅黑")
    normal.font.size = Pt(9.25)

    p = doc.add_paragraph()
    paragraph_base(p, after=1)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(p.add_run(payload["name"]), size=18, name="Arial")
    p = doc.add_paragraph()
    paragraph_base(p, after=5)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(p.add_run(payload.get("contact", "")), size=8.5, name="Arial")

    if payload.get("summary"):
        section_title(doc, "个人简介")
        p = doc.add_paragraph()
        paragraph_base(p, after=2, line=1.05, keep=True)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        font(p.add_run(payload["summary"]), size=9.3)

    education = payload.get("education", [])
    experience = list(payload.get("experience", []))
    projects = list(payload.get("projects", []))
    if projects_position == "integrated":
        experience += projects
        experience.sort(key=lambda item: item.get("sort_date", ""), reverse=True)
        projects = []

    if education_position == "before_experience":
        add_education(doc, education)
    if projects_position == "before_experience":
        add_entries(doc, "项目经历", projects)
    add_entries(doc, "工作经历", experience)
    if projects_position == "after_experience":
        add_entries(doc, "项目经历", projects)
    add_entries(doc, "实习经历", payload.get("internships", []))
    if education_position == "after_experience":
        add_education(doc, education)
    add_skills(doc, payload.get("skills", []))
    doc.save(path)


def make_docm(docx_path, output_path):
    with zipfile.ZipFile(docx_path) as source, zipfile.ZipFile(output_path, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(
                    b"application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml",
                    b"application/vnd.ms-word.document.macroEnabled.main+xml",
                )
            target.writestr(item, data)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    if args.output.suffix.lower() not in {".docx", ".docm"}:
        raise SystemExit("Output must end in .docx or .docm")
    payload = json.loads(args.input.read_text(encoding="utf-8"))
    args.output.parent.mkdir(parents=True, exist_ok=True)
    if args.output.suffix.lower() == ".docx":
        write_docx(payload, args.output)
    else:
        with tempfile.TemporaryDirectory() as temp:
            source = Path(temp) / "source.docx"
            write_docx(payload, source)
            make_docm(source, args.output)
    print(args.output)


if __name__ == "__main__":
    main()
